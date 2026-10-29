"""Only web and configured local directory enter the evidence pipeline."""

import asyncio
import csv
import hashlib
import io
import ipaddress
import json
import re
import socket
from bisect import bisect_right
from functools import lru_cache
from pathlib import Path
from urllib.parse import urlparse
from typing import Protocol
import httpx
from bs4 import BeautifulSoup
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from .schemas import SourceRecord, uid
from .storage import Store

MAX_BYTES = 20 * 1024 * 1024
EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.csv', '.txt', '.md', '.html', '.htm', '.json'}


class DatasourceError(RuntimeError):
    pass


NUMBERED_PART = re.compile(r'^([a-z]+):(\d+)$')
CHAR_PART = re.compile(r'^chars:(\d+)-(\d+)$')
CHUNK_MARKER = re.compile(r'^\[([a-z]+:[^\]\n]+)\]$', re.M)


def compact_locator(locator: str | None) -> str | None:
    """'page:1; page:2; page:3 chars 10:20' -> 'page:1-3 chars 10:20'.

    Joined documents list every chunk; consecutive pages or character windows
    collapse into one range and anything unrecognised is kept as written."""
    if not locator:
        return locator
    head, sep, tail = locator.partition(' chars ')
    out = []
    run = None  # (kind, first, last) for numbered parts, ('chars', start, end) for windows

    def flush():
        if run:
            kind, first, last = run
            if kind == 'chars':
                out.append(f'chars:{first}-{last}')
            else:
                out.append(f'{kind}:{first}' + (f'-{last}' if last != first else ''))

    for part in (p.strip() for p in head.split(';')):
        if not part:
            continue
        numbered, chars = NUMBERED_PART.match(part), CHAR_PART.match(part)
        if numbered and run and run[0] == numbered[1] and int(numbered[2]) == run[2] + 1:
            run = (run[0], run[1], run[2] + 1)
            continue
        if chars and run and run[0] == 'chars' and int(chars[1]) == run[2]:
            run = ('chars', run[1], int(chars[2]))
            continue
        flush()
        run = None
        if numbered:
            run = (numbered[1], int(numbered[2]), int(numbered[2]))
        elif chars:
            run = ('chars', int(chars[1]), int(chars[2]))
        else:
            out.append(part)
    flush()
    return '; '.join(out) + sep + tail


@lru_cache(maxsize=64)
def chunk_markers(content: str) -> tuple[tuple[int, str], ...]:
    return tuple((m.start(), m.group(1)) for m in CHUNK_MARKER.finditer(content))


def span_locator(source: SourceRecord, start: int, end: int) -> str:
    """Locator of the chunks (pages, rows) a quote at content[start:end] actually spans."""
    markers = chunk_markers(source.content)
    first = bisect_right([pos for pos, _ in markers], start) - 1
    labels = [label for pos, label in markers[max(first, 0) :] if pos < end] if first >= 0 else []
    where = compact_locator('; '.join(labels)) if labels else compact_locator(source.locator)
    return f'{where} chars {start}:{end}'


def quote_locator(source: SourceRecord, quote: str) -> str:
    start = source.content.find(quote) if quote else -1
    return span_locator(source, start, start + len(quote)) if start >= 0 else compact_locator(source.locator)


def clean_html(text: str) -> str:
    soup = BeautifulSoup(text, 'html.parser')
    for tag in soup(['script', 'style', 'iframe', 'object', 'form', 'noscript']):
        tag.decompose()
    return soup.get_text(' ', strip=True)


def parse_material(data: bytes, suffix: str) -> list[tuple[str, str]]:
    """Return located chunks. PDF without text needs OCR; never silently invented."""
    if len(data) > MAX_BYTES:
        raise DatasourceError('材料超过大小上限')
    if suffix == '.pdf':
        chunks = [(f'page:{i + 1}', p.extract_text() or '') for i, p in enumerate(PdfReader(io.BytesIO(data)).pages)]
        if not any(t.strip() for _, t in chunks):
            raise DatasourceError('扫描 PDF 无可提取文本，需提供 OCR 后材料')
        return chunks
    if suffix == '.docx':
        doc = Document(io.BytesIO(data))
        chunks = []
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                chunks.append((f'paragraph:{i + 1}', p.text))
        for i, table in enumerate(doc.tables):
            for j, row in enumerate(table.rows):
                chunks.append((f'table:{i + 1}/row:{j + 1}', ' | '.join(c.text for c in row.cells)))
        return chunks
    if suffix == '.xlsx':
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        chunks = []
        try:
            for sheet in workbook:
                headers = None
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if not any(c is not None for c in row):
                        continue
                    names = {str(c) for c in row}
                    if i == 0 and any(
                        required.issubset(names)
                        for required in [
                            {'company_id', 'metric_code', 'period', 'unit', 'value', 'actual_or_forecast'},
                            {'company_id', 'field_id', 'period', 'unit', 'value'},
                            {'company_id', 'ticker', 'company_name'},
                        ]
                    ):
                        headers = [str(c) for c in row]
                        continue
                    text = (
                        json.dumps(dict(zip(headers, row)), ensure_ascii=False, default=str)
                        if headers
                        else ' | '.join('' if c is None else str(c) for c in row)
                    )
                    chunks.append((f'sheet:{sheet.title}/row:{i + 1}/cells:A-{len(row)}', text))
        finally:
            workbook.close()
        return chunks
    text = data.decode('utf-8-sig', errors='replace')
    if suffix == '.csv':
        reader = csv.DictReader(io.StringIO(text))
        return [(f'row:{i + 2}', json.dumps(row, ensure_ascii=False)) for i, row in enumerate(reader)]
    if suffix == '.json':
        body = json.loads(text)
        rows = body if isinstance(body, list) else body.get('facts', [body])
        if isinstance(body, dict) and 'facts' in body:
            metadata = {
                k: body[k]
                for k in (
                    'document_type',
                    'company_id',
                    'company_name',
                    'ticker',
                    'security_name',
                    'market',
                    'financial_company',
                    'report_period',
                    'currency',
                    'accounting_basis',
                    'segments',
                    'industries',
                    'publication_date',
                    'period_kind',
                )
                if k in body
            }
            rows = [metadata | row if isinstance(row, dict) else row for row in rows]
        return [(f'record:{i + 1}', json.dumps(row, ensure_ascii=False)) for i, row in enumerate(rows)]
    if suffix in {'.html', '.htm'}:
        text = clean_html(text)
    # Limited chunk sizes preserve retrievable location and avoid unbounded context.
    return [(f'chars:{i}-{min(i + 5000, len(text))}', text[i : i + 5000]) for i in range(0, len(text), 5000)]


def source_record(kind: str, title: str, content: str, locator: str, **kwargs) -> SourceRecord:
    digest = hashlib.sha256(
        (str(kwargs.get('local_path') or kwargs.get('source_url')) + locator + content).encode()
    ).hexdigest()
    return SourceRecord(
        source_id=digest[:20],
        kind=kind,
        title=title,
        content=content,
        locator=locator,
        content_hash=hashlib.sha256(content.encode()).hexdigest(),
        **kwargs,
    )


def join_document_chunks(records):
    """Preserve page/table/row boundaries while making preceding headers available across chunks."""
    groups = {}
    structured = []
    for s in records:
        if s.content.lstrip().startswith('{'):
            structured.append(s)
            continue
        groups.setdefault((s.kind, s.title, s.local_path, s.source_url), []).append(s)
    for (_, _, _, _), chunks in groups.items():
        if len(chunks) == 1:
            structured += chunks
            continue
        first = chunks[0]
        content = '\n\n'.join('[' + s.locator + ']\n' + s.content for s in chunks)
        joined = '; '.join(s.locator for s in chunks)
        record = source_record(
            first.kind,
            first.title,
            content,
            joined,
            local_path=first.local_path,
            source_url=first.source_url,
            publication_date=first.publication_date,
            attachment_id=first.attachment_id,
        )
        # The ID still hashes the full chunk list so earlier runs keep resolving;
        # only the stored locator is shortened (page:1; page:2; ... -> page:1-124).
        structured.append(record.model_copy(update={'locator': compact_locator(joined)}))
    return structured


class SearchProvider(Protocol):
    async def search(self, query: str, limit: int = 6) -> list[dict]: ...


class HTTPWebSearchProvider:
    """Injected search endpoint: POST {query,limit} -> {results:[{url,title,...}]}.

    A vendor's search payload can be adapted here without adding a data source.
    """

    def __init__(self, url: str, key: str = '', auth_header='Authorization', client=None):
        self.url, self.key, self.auth_header, self.client = url, key, auth_header, client

    async def search(self, query: str, limit: int = 6):
        if not self.url:
            raise DatasourceError('Web Search 未配置 CDR_SEARCH_URL')
        headers = {self.auth_header: self.key} if self.key else {}
        for attempt in range(3):
            try:
                if self.client:
                    response = await self.client.post(
                        self.url, json={'query': query, 'limit': limit}, headers=headers, timeout=30
                    )
                else:
                    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
                        response = await client.post(self.url, json={'query': query, 'limit': limit}, headers=headers)
                response.raise_for_status()
                body = response.json()
                if not isinstance(body.get('results'), list):
                    raise DatasourceError('搜索结果不符合已配置 JSON 协议')
                return body['results'][:limit]
            except (httpx.HTTPError, ValueError):
                if attempt == 2:
                    raise DatasourceError('Web Search 请求失败') from None
                await asyncio.sleep((1, 2, 4)[attempt])
        return []


class BochaWebSearchProvider:
    """Adapter for Bocha's Web Search API while retaining the common result shape."""

    def __init__(self, url: str, key: str, client=None):
        self.url, self.key, self.client = url, key, client

    async def search(self, query: str, limit: int = 6) -> list[dict]:
        if not self.url:
            raise DatasourceError('Web Search 未配置 CDR_SEARCH_URL')
        if not self.key:
            raise DatasourceError('博查搜索未配置 CDR_SEARCH_KEY_FILE 或 CDR_SEARCH_KEY')
        headers = {'Authorization': f'Bearer {self.key}'}
        payload = {'query': query, 'count': limit, 'summary': False}
        for attempt in range(3):
            try:
                if self.client:
                    response = await self.client.post(self.url, json=payload, headers=headers, timeout=30)
                else:
                    async with httpx.AsyncClient(timeout=30, trust_env=False) as client:
                        response = await client.post(self.url, json=payload, headers=headers)
                response.raise_for_status()
                body = response.json()
                data = body.get('data', body)
                pages = data.get('webPages', {}).get('value', [])
                if not isinstance(pages, list):
                    raise DatasourceError('博查搜索结果不符合预期格式')
                return [
                    {
                        'url': item['url'],
                        'title': item.get('name', item['url']),
                        'publication_date': item.get('datePublished') or item.get('dateLastCrawled'),
                    }
                    for item in pages[:limit]
                    if isinstance(item, dict) and isinstance(item.get('url'), str)
                ]
            except (httpx.HTTPError, ValueError, KeyError):
                if attempt == 2:
                    raise DatasourceError('博查 Web Search 请求失败') from None
                await asyncio.sleep((1, 2, 4)[attempt])
        return []


class LocalDatasource:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.index: list[SourceRecord] = []
        self.errors: list[str] = []

    def build_index(self, exclude_annual_reports=False) -> None:
        from .annual_reports import annual_filename
        from .reporting_periods import financial_document, INTERIM

        self.index = []
        self.errors = []
        if not self.root.exists():
            return
        for path in sorted(self.root.rglob('*')):
            if not path.is_file() or path.suffix.lower() not in EXTENSIONS:
                continue
            if exclude_annual_reports and (annual_filename(path.name) or re.search(INTERIM, path.name, re.I)):
                continue
            if not path.resolve().is_relative_to(self.root) or path.is_symlink():
                self.errors.append('本地来源越过指定目录，已拒绝')
                continue
            try:
                if path.stat().st_size > MAX_BYTES:
                    raise DatasourceError('材料超过大小上限')
                for locator, text in parse_material(path.read_bytes(), path.suffix.lower()):
                    if text.strip():
                        source = source_record(
                            'local', path.name, text, locator, local_path=str(path.relative_to(self.root))
                        )
                        if not exclude_annual_reports or not financial_document(source):
                            self.index.append(source)
            except Exception:
                self.errors.append(f'{path.name} 解析失败或需 OCR，请检查材料格式')

    def search(self, query: str, limit: int = 8) -> list[SourceRecord]:
        terms = re.findall(r'[\w]+', query.lower())
        cjk = re.findall(r'[\u4e00-\u9fff]{2,}', query)
        terms += [t[i : i + 2] for t in cjk for i in range(len(t) - 1)]
        scored = [(sum(term in (s.title + ' ' + s.content).lower() for term in terms), s) for s in self.index]
        selected = [s for score, s in sorted(scored, key=lambda p: p[0], reverse=True)[:limit] if score > 0]
        documents = {s.local_path for s in selected if not s.content.lstrip().startswith('{')}
        rows = [s for s in selected if s.content.lstrip().startswith('{')]
        return rows + join_document_chunks(
            [s for s in self.index if s.local_path in documents and not s.content.lstrip().startswith('{')]
        )

    def exact_rows(self, **fields) -> list[SourceRecord]:
        result = []
        for s in self.index:
            try:
                row = json.loads(s.content)
            except ValueError:
                continue
            if all(str(row.get(k, '')) == str(v) for k, v in fields.items()):
                result.append(s)
        return result


async def validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {'https', 'http'} or not parsed.hostname or parsed.username or parsed.password:
        raise DatasourceError('只允许公开 HTTP/HTTPS 来源 URL')
    if parsed.port not in {None, 80, 443}:
        raise DatasourceError('拒绝非公开网页端口')
    try:
        addresses = await asyncio.to_thread(
            socket.getaddrinfo, parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM
        )
        if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
            raise DatasourceError('禁止访问私有网络或本机地址')
    except socket.gaierror:
        raise DatasourceError('来源地址解析失败') from None


class WebAdapter:
    def __init__(self, provider: SearchProvider, store: Store, client=None, *, run_id=None):
        self.provider, self.store, self.client = provider, store, client
        self.cache_scope = run_id or uid()
        self.read_limit = asyncio.Semaphore(3)

    async def search(self, query: str) -> list[dict]:
        key = f'run:{self.cache_scope}:search:' + hashlib.sha256(query.encode()).hexdigest()
        cached = self.store.cached(key)
        if cached is not None:
            return cached
        data = await self.provider.search(query)
        self.store.cache(key, data)
        return data

    async def read(self, result: dict) -> list[SourceRecord]:
        url = result['url']
        key = f'run:{self.cache_scope}:read:' + hashlib.sha256(url.encode()).hexdigest()
        cached = self.store.cached(key)
        if cached:
            return [SourceRecord.model_validate(s) for s in cached]
        async with self.read_limit:
            await validate_public_url(url)
            for attempt in range(3):
                try:
                    async with httpx.AsyncClient(timeout=30, trust_env=False, follow_redirects=False) as owned:
                        client = self.client or owned
                        current = url
                        for _ in range(5):
                            await validate_public_url(current)
                            async with client.stream(
                                'GET', current, headers={'User-Agent': 'CompanyResearch/0.1 public-document-reader'}
                            ) as response:
                                if response.is_redirect:
                                    from urllib.parse import urljoin

                                    current = urljoin(current, response.headers['location'])
                                    continue
                                response.raise_for_status()
                                chunks = []
                                size = 0
                                async for chunk in response.aiter_bytes():
                                    size += len(chunk)
                                    if size > MAX_BYTES:
                                        raise DatasourceError('公开文件超过下载上限')
                                    chunks.append(chunk)
                                data = b''.join(chunks)
                                ctype = response.headers.get('content-type', '')
                                break
                        else:
                            raise DatasourceError('来源重定向次数过多')
                    suffix = Path(urlparse(current).path).suffix.lower()
                    if 'pdf' in ctype:
                        suffix = '.pdf'
                    if suffix not in EXTENSIONS:
                        suffix = '.html'
                    sources = join_document_chunks(
                        [
                            source_record(
                                'web',
                                result.get('title', url),
                                text,
                                locator,
                                source_url=current,
                                publication_date=result.get('publication_date'),
                            )
                            for locator, text in parse_material(data, suffix)
                            if text.strip()
                        ]
                    )
                    self.store.cache(key, [s.model_dump() for s in sources])
                    return sources
                except httpx.HTTPError:
                    if attempt == 2:
                        raise DatasourceError('公开页面或文档读取失败') from None
                    await asyncio.sleep(0.1 * 2**attempt)
        return []


def domain_match(host: str, domain: str) -> bool:
    return host == domain or host.endswith('.' + domain)


class SourcePolicy:
    """Domain table from configs/source_domains.yaml; filtering happens before any read."""

    def __init__(self, blocked_domains=(), blocked_url_patterns=(), preferred_domains=(), policy_site=''):
        self.blocked_domains = [d.lower().strip('.') for d in blocked_domains]
        self.blocked_patterns = [re.compile(p, re.I) for p in blocked_url_patterns]
        self.preferred_domains = [d.lower().strip('.') for d in preferred_domains]
        self.policy_site = (policy_site or '').strip()

    @classmethod
    def load(cls, path: Path | None) -> 'SourcePolicy':
        if not path or not Path(path).is_file():
            return cls()
        import yaml

        data = yaml.safe_load(Path(path).read_text(encoding='utf-8')) or {}
        return cls(
            data.get('blocked_domains') or [],
            data.get('blocked_url_patterns') or [],
            data.get('preferred_domains') or [],
            data.get('policy_site') or '',
        )

    def blocked(self, url: str) -> bool:
        host = (urlparse(url).hostname or '').lower()
        return any(domain_match(host, d) for d in self.blocked_domains) or any(
            p.search(url) for p in self.blocked_patterns
        )

    def rank(self, url: str) -> int:
        host = (urlparse(url).hostname or '').lower()
        return next(
            (i for i, d in enumerate(self.preferred_domains) if domain_match(host, d)), len(self.preferred_domains)
        )

    def select(self, results: list[dict]) -> list[dict]:
        kept = [r for r in results if isinstance(r.get('url'), str) and not self.blocked(r['url'])]
        return sorted(kept, key=lambda r: self.rank(r['url']))  # Stable: unlisted keep engine order.


class UnifiedDatasource:
    def __init__(self, local: LocalDatasource, web: WebAdapter | None, policy: SourcePolicy | None = None):
        self.local, self.web = local, web
        self.policy = policy or SourcePolicy()
        self.exclude_annual_reports = False

    async def retrieve(self, query: str, web_limit=4, filter_sources=False) -> tuple[list[SourceRecord], list[str]]:
        from .annual_reports import annual_search_result, SECONDARY
        from .reporting_periods import financial_document, INTERIM

        records = self.local.search(query)
        warnings = []
        if self.exclude_annual_reports:
            records = [s for s in records if not financial_document(s)]
        if self.web:
            results = await self.web.search(query)
            if self.exclude_annual_reports:
                results = [
                    r
                    for r in results
                    if not annual_search_result(r)
                    and not (
                        re.search(
                            INTERIM, ' '.join(str(r.get(k) or '') for k in ('title', 'description', 'snippet')), re.I
                        )
                        and not re.search(SECONDARY, r.get('title', ''), re.I)
                    )
                ]
            if filter_sources:
                # Blocked previews never consume the read budget; the next result is read instead.
                results = self.policy.select(results)
            read = await asyncio.gather(*(self.web.read(r) for r in results[:web_limit]), return_exceptions=True)
            for r in read:
                if isinstance(r, Exception):
                    warnings.append('一项公开来源读取失败，已保留缺口')
                else:
                    records.extend(s for s in r if not self.exclude_annual_reports or not financial_document(s))
        unique = {s.source_id: s for s in records}
        return list(unique.values()), warnings + self.local.errors
