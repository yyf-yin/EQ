"""Safe structured rendering; document numbers always originate in Python."""

from html import escape
import re
from pathlib import Path
from .schemas import ReportDocument, ReportSection, Table, SourceRecord

TITLES = [
    "核心观点",
    "目录",
    "一、公司简介",
    "二、财务情况分析",
    "三、公司业务",
    "四、盈利预测与估值",
    "五、风险提示",
    "六、财务三表",
    "七、免责声明",
]


def sections_flat(sections: list[ReportSection]):
    for section in sections:
        yield section
        yield from sections_flat(section.children)


def toc(sections: list[ReportSection]) -> list[tuple[str, str]]:
    return [(s.section_id, s.title) for s in sections_flat(sections) if s.section_id not in {"core", "toc"}]


def mock_report(company: str) -> ReportDocument:
    source = SourceRecord(
        source_id="mock-fixture",
        kind="mock",
        title="模拟财务与经营材料",
        content="所有数据均为程序测试合成，非真实公司资料。",
    )
    sections = [
        ReportSection(section_id=sid, title=title, paragraphs=[text])
        for sid, title, text in [
            (
                "core",
                TITLES[0],
                "本报告为模拟示例，仅验证研究、计算、审核和导出流程。演示公司收入增长与现金创造能力的关系，所有数字均为合成数据，不可用于投资判断。",
            ),
            ("toc", TITLES[1], ""),
            (
                "company",
                TITLES[2],
                f"输入对象为{company}。模拟模式使用虚构的示例制造企业，包含设备和服务两个业务。输入名称不会使模拟材料成为真实资料。",
            ),
            (
                "financial",
                TITLES[3],
                "模拟企业收入增加时，应同时关注应收与库存增长。利润和现金流可能不同步；营运资本占用影响经营现金流，扩产支出影响自由现金流。",
            ),
            ("business", TITLES[4], "从设备与服务两个业务分别分析需求、竞争和经营质量。"),
            (
                "forecast",
                TITLES[5],
                "预测采用公开的示例假设。情景变化由 Python 确定性计算，真实研究需要验证需求、价格、成本和资本开支的来源。",
            ),
            (
                "risk",
                TITLES[6],
                "主要风险为需求波动、客户回款延迟、成本上涨和资本开支超支。模拟数据不反映输入公司的真实风险。",
            ),
            ("statements", TITLES[7], "此阶段仅提供测试报表；完整财务模型由后续计算底座接入。"),
            (
                "disclaimer",
                TITLES[8],
                "本系统生成研究辅助材料，不构成投资建议。模拟报告全部使用虚构资料。预测属于条件推演，不能视为实际业绩或收益保证。",
            ),
        ]
    ]
    for index, name in enumerate(["设备", "服务"]):
        sections[4].children.append(
            ReportSection(
                section_id=f"segment-{index}",
                title=name,
                children=[
                    ReportSection(
                        section_id=f"industry-{index}",
                        title="（一）行业情况",
                        paragraphs=["演示行业分析应识别需求周期、技术替代与竞争结构。真实判断必须依赖可核查证据。"],
                    ),
                    ReportSection(
                        section_id=f"position-{index}",
                        title="（二）公司情况",
                        paragraphs=[
                            "演示经营分析应将客户结构、产能约束和盈利质量联系起来，不将缺少证据的优势写成确定事实。"
                        ],
                    ),
                ],
            )
        )
    sections[7].tables.append(
        Table(
            title="模拟利润表 单位 百万元",
            columns=["项目", "2024A", "2025A"],
            rows=[["营业收入", "1000", "1100"], ["营业成本", "600", "660"]],
        )
    )
    return ReportDocument(
        title=f"{company} 公司深度研究报告",
        company_id="MOCK",
        mode="mock",
        sections=sections,
        sources=[source],
        limitations=["模拟运行：全部材料为合成 Fixture，未调用真实模型或搜索。"],
    )


NUMERIC = re.compile(r'^[-+—]?[\d,]*\.?\d*%?$|^—$')


def html_report(report: ReportDocument) -> str:
    from .charts import to_svg
    from .presentation import lead, present

    view = present(report)

    def para(text):
        head, rest = lead(text)
        return f'<p>{"<b>" + escape(head) + "</b>" if head else ""}{escape(rest)}</p>'

    def table(t, numbered=True):
        caption = (f'<b>{escape(t.label)}：</b>' if numbered and t.label else '') + escape(t.title)
        rows = ''
        for i, r in enumerate(t.rows):
            if i in t.groups:
                rows += f'<tr class="group"><td colspan="{len(t.columns)}">{escape(r[0])}</td></tr>'
                continue
            rows += (
                '<tr>'
                + ''.join(f'<td class="{"num" if j and NUMERIC.match(c) else ""}">{escape(c)}</td>' for j, c in enumerate(r))
                + '</tr>'
            )
        head = ''.join(f'<th>{escape(c)}</th>' for c in t.columns)
        return (
            f'<figure class="tbl"><figcaption>{caption}</figcaption><div class="scroll"><table><thead><tr>{head}</tr></thead>'
            f'<tbody>{rows}</tbody></table></div><p class="src">数据来源：{escape(t.source)}</p>'
            + (f'<p class="src">注：{escape(t.note)}</p>' if t.note else '')
            + '</figure>'
        )

    def chart(c):
        return (
            f'<figure class="fig"><figcaption><b>{escape(c.label)}：</b>{escape(c.title)}</figcaption>'
            f'{to_svg(c)}<p class="src">数据来源：{escape(c.source)}</p></figure>'
        )

    def blocks(items):
        out, row = [], []
        for kind, item in items + [('end', None)]:
            if kind == 'chart':
                row.append(chart(item))
                continue
            if row:
                out.append('<div class="figs">' + ''.join(row) + '</div>')
                row = []
            if kind == 'p':
                out.append(para(item))
            elif kind == 'table':
                out.append(table(item))
        return ''.join(out)

    def section(s):
        tag = f'h{min(s.level + 1, 4)}'
        return (
            f'<section id="{escape(s.section_id, quote=True)}"><{tag}>{escape(s.title)}</{tag}>'
            + blocks(s.blocks)
            + ''.join(section(c) for c in s.children)
            + '</section>'
        )

    points = ''.join(f'<p>{"<b>" + escape(h) + "</b>" if h else ""}{escape(b)}</p>' for h, b in view.points)
    summary = table(view.forecast_table, numbered=False) if view.forecast_table else ''
    contents = ''.join(
        f'<li class="l{s.level}"><a href="#{escape(s.section_id, quote=True)}">{escape(s.title)}</a></li>'
        for s in view.flat()
    ) + ('<li class="l1"><a href="#appendix">附录：财务报表</a></li>' if view.appendix else '')
    figures = ''.join(f'<li>{escape(f.label)}：{escape(f.title)}</li>' for f in view.figures())
    appendix = (
        '<section id="appendix"><h2>附录：财务报表</h2>' + ''.join(table(t, False) for t in view.appendix) + '</section>'
        if view.appendix
        else ''
    )
    body = (
        f'<header class="cover"><p class="kicker">{escape(view.subtitle)}</p><h1>{escape(view.title)}</h1>'
        f'<div class="cover-grid"><div><h2 class="core">核心观点</h2>{points}</div>'
        f'<aside>{summary}</aside></div></header>'
        f'<nav class="toc" aria-label="目录"><h2>目录</h2><ol>{contents}</ol>'
        + (f'<h3>图表索引</h3><ol class="figlist">{figures}</ol>' if figures else '')
        + '</nav>'
        + ''.join(section(s) for s in view.sections if s.blocks or s.children)
        + appendix
        + f'<footer class="disclaimer"><b>免责声明：</b>{escape(view.disclaimer)}</footer>'
    )
    return (
        '<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1"><title>'
        + escape(view.title)
        + '</title><style>'
        + REPORT_CSS
        + '</style><body><main>'
        + body
        + '</main></body></html>'
    )


REPORT_CSS = """
:root{--ink:#222;--muted:#666;--brand:#8c1d18;--line:#d9d9d9;--head:#f3f0ef;--paper:#fff;--bg:#f5f5f4}
*{box-sizing:border-box}html{background:var(--bg)}body{margin:0;background:var(--bg);color:var(--ink);
font:15px/1.85 "Microsoft YaHei","PingFang SC","Noto Sans CJK SC",system-ui,sans-serif}
main{max-width:980px;margin:24px auto;background:var(--paper);padding:40px 48px;box-shadow:0 1px 3px rgba(0,0,0,.08)}
h1{font-size:26px;margin:4px 0 18px;color:var(--brand)}
h2{font-size:20px;color:var(--brand);border-bottom:2px solid var(--brand);padding-bottom:6px;margin:40px 0 14px}
h3{font-size:16.5px;margin:28px 0 8px}h4{font-size:15px;margin:20px 0 6px}p{margin:0 0 12px;text-align:justify}
.kicker{color:var(--muted);letter-spacing:.1em;margin:0}.cover{border-bottom:1px solid var(--line);padding-bottom:16px}
.cover-grid{display:grid;grid-template-columns:minmax(0,1.6fr) minmax(0,1fr);gap:28px}.cover h2.core{margin-top:0}
.toc ol{list-style:none;padding:0;margin:0}.toc li.l2{padding-left:1.5em}.toc li.l3{padding-left:3em;font-size:14px}
.toc a{color:var(--ink);text-decoration:none}.toc a:hover{color:var(--brand)}
.figlist{columns:2;font-size:13px;color:var(--muted)}
.figs{display:grid;grid-template-columns:repeat(auto-fit,minmax(300px,1fr));gap:18px;margin:8px 0 18px}
figure{margin:0}figcaption{font-size:14px;margin-bottom:4px}figure svg{width:100%;height:auto;display:block}
.src{font-size:12px;color:var(--muted);margin:2px 0 0}.tbl{margin:10px 0 20px}.scroll{overflow-x:auto}
table{border-collapse:collapse;width:100%;font-size:13px;border-top:2px solid var(--brand)}
th{background:var(--head);font-weight:600;padding:5px 8px;text-align:center;white-space:nowrap}
td{padding:4px 8px;border-bottom:1px solid var(--line)}
td.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
tr.group td{font-weight:600;background:#faf8f7}aside .tbl{margin-top:6px}
.disclaimer{margin-top:40px;font-size:12px;color:var(--muted);border-top:1px solid var(--line);padding-top:10px}
@media (max-width:900px){.cover-grid{grid-template-columns:1fr}}
@media (max-width:760px){main{padding:20px 16px;margin:0}.figlist{columns:1}}
"""


def export_docx(report: ReportDocument, path: Path) -> None:
    import io

    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    from .charts import to_png
    from .presentation import lead, present

    view = present(report)
    brand = RGBColor(0x8C, 0x1D, 0x18)
    muted = RGBColor(0x66, 0x66, 0x66)
    document = Document()
    sec = document.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2)
    width = sec.page_width - sec.left_margin - sec.right_margin
    for style in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4']:
        st = document.styles[style]
        st.font.name = 'Microsoft YaHei'
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.color.rgb = brand if style in ('Title', 'Heading 1') else RGBColor(0x22, 0x22, 0x22)
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    for name, size in [('Title', 20), ('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11)]:
        document.styles[name].font.size = Pt(size)
        document.styles[name].paragraph_format.keep_with_next = True

    def small(paragraph, size=8, color=muted):
        for r in paragraph.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = color
        paragraph.paragraph_format.space_after = Pt(8)
        return paragraph

    def para(text):
        head, rest = lead(text)
        p = document.add_paragraph()
        if head:
            p.add_run(head).bold = True
        p.add_run(rest)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def caption(paragraph, text, label=''):
        if label:
            paragraph.add_run(label + '：').bold = True
        paragraph.add_run(text).bold = True
        for r in paragraph.runs:
            r.font.size = Pt(9)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(2)
        return paragraph

    def borders(table):
        element = OxmlElement('w:tblBorders')
        for edge, size, color in [('top', 12, '8C1D18'), ('bottom', 6, '999999'), ('insideH', 2, 'D9D9D9')]:
            item = OxmlElement(f'w:{edge}')
            item.set(qn('w:val'), 'single')
            item.set(qn('w:sz'), str(size))
            item.set(qn('w:color'), color)
            element.append(item)
        table._tbl.tblPr.append(element)

    def data_table(t, numbered=True):
        caption(document.add_paragraph(), t.title, t.label if numbered else '')
        table = document.add_table(rows=1, cols=len(t.columns))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(table)
        first = 0.3 if len(t.columns) > 3 else 0.4
        widths = [int(width * first)] + [int(width * (1 - first) / (len(t.columns) - 1))] * (len(t.columns) - 1)
        for i, value in enumerate(t.columns):
            cell = table.rows[0].cells[i]
            cell.text = value
            shade = OxmlElement('w:shd')
            shade.set(qn('w:val'), 'clear')
            shade.set(qn('w:fill'), 'F3F0EF')
            cell._tc.get_or_add_tcPr().append(shade)
            cell.paragraphs[0].runs[0].bold = True
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for index, row in enumerate(t.rows):
            cells = table.add_row().cells
            if index in t.groups:
                merged = cells[0].merge(cells[-1])
                merged.text = row[0]
                merged.paragraphs[0].runs[0].bold = True
                continue
            for j, value in enumerate(row):
                cells[j].text = value
                if j and NUMERIC.match(value):
                    cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for i, cell in enumerate(row.cells):
                cell.width = widths[min(i, len(widths) - 1)]
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
        small(document.add_paragraph('数据来源：' + t.source + (f'\n注：{t.note}' if t.note else '')))

    def figure_row(items):
        """Charts two per row, caption above and source below each, as broker reports lay them out."""
        for start in range(0, len(items), 2):
            pair = items[start : start + 2]
            grid = document.add_table(rows=1, cols=len(pair))
            grid.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, c in zip(grid.rows[0].cells, pair):
                cell.width = int(width / len(pair))
                caption(cell.paragraphs[0], c.title, c.label)
                stream = io.BytesIO()
                to_png(c, stream)
                stream.seek(0)
                cell.add_paragraph().add_run().add_picture(stream, width=int(width / len(pair) * 0.95))
                small(cell.add_paragraph('数据来源：' + c.source), 7.5)
            document.add_paragraph().paragraph_format.space_after = Pt(0)

    def blocks(items):
        row = []
        for kind, item in items + [('end', None)]:
            if kind == 'chart':
                row.append(item)
                continue
            if row:
                figure_row(row)
                row = []
            if kind == 'p':
                para(item)
            elif kind == 'table':
                data_table(item)

    small(document.add_paragraph(view.subtitle), 10)
    document.add_heading(view.title, 0)
    head = document.add_paragraph()
    run = head.add_run('核心观点：')
    run.bold = True
    run.font.color.rgb = brand
    run.font.size = Pt(12)
    for h, b in view.points:
        p = document.add_paragraph()
        if h:
            p.add_run(h).bold = True
        p.add_run(b)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if view.forecast_table:
        data_table(view.forecast_table, numbered=False)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    document.add_heading('目录索引', 1)
    for s in view.flat():
        p = document.add_paragraph(s.title)
        p.paragraph_format.left_indent = Cm(0.8 * (s.level - 1))
        p.paragraph_format.space_after = Pt(2)
    if view.appendix:
        document.add_paragraph('附录：财务报表').paragraph_format.space_after = Pt(2)
    figures = view.figures()
    if figures:
        document.add_heading('图表索引', 1)
        for f in figures:
            small(document.add_paragraph(f'{f.label}：{f.title}'), 9, RGBColor(0x33, 0x33, 0x33)).paragraph_format.space_after = Pt(1)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def section(s):
        document.add_heading(s.title, min(s.level, 3))
        blocks(s.blocks)
        for child in s.children:
            section(child)

    for s in view.sections:
        if s.blocks or s.children:
            section(s)
    if view.appendix:
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_heading('附录：财务报表', 1)
        for t in view.appendix:
            data_table(t, numbered=False)
    small(document.add_paragraph('免责声明：' + view.disclaimer), 8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
