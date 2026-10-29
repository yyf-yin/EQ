import csv,io,json
from pathlib import Path
import httpx
import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from company_research.datasources import LocalDatasource,HTTPWebSearchProvider,WebAdapter,parse_material,validate_public_url,DatasourceError
from company_research.evidence import EvidenceStore
from company_research.storage import Store


def test_all_formats_and_exact(tmp_path):
    (tmp_path/'a.txt').write_text('设备 行业资料',encoding='utf-8')
    (tmp_path/'a.md').write_text('# 设备行业',encoding='utf-8')
    (tmp_path/'a.html').write_text('<h1>设备</h1><script>evil</script>',encoding='utf-8')
    row={'company_id':'c','metric_code':'revenue','period':'2025A','unit':'CNY_million','value':'100','actual_or_forecast':'actual'}
    (tmp_path/'a.json').write_text(json.dumps([row]),encoding='utf-8')
    with (tmp_path/'a.csv').open('w',newline='',encoding='utf-8') as f:
        writer=csv.DictWriter(f,fieldnames=row.keys());writer.writeheader();writer.writerow(row)
    doc=Document();doc.add_paragraph('设备公司');doc.add_table(rows=1,cols=1).cell(0,0).text='设备表格';doc.save(tmp_path/'a.docx')
    wb=Workbook();wb.active.append(['设备','100']);wb.save(tmp_path/'a.xlsx')
    pdf=canvas.Canvas(str(tmp_path/'a.pdf'));pdf.drawString(50,700,'Company equipment');pdf.save()
    local=LocalDatasource(tmp_path);local.build_index()
    assert not local.errors
    assert len(local.index)>=8
    assert local.search('设备')
    exact=local.exact_rows(company_id='c',metric_code='revenue',period='2025A');assert len(exact)==2
    store=EvidenceStore(Store(tmp_path/'store.sqlite'),'run')
    for source in exact:store.ingest_structured(source)
    assert len(store.facts)==2 and all(f.value==100 for f in store.facts.values())
    assert 'evil' not in next(s.content for s in local.index if s.title=='a.html')
    assert any(s.locator.startswith('page:') for s in local.index)
    assert any(s.locator.startswith('table:') for s in local.index)


async def test_web_cache_retry_and_config(tmp_path):
    calls=0
    def handler(request):
        nonlocal calls
        calls+=1
        if calls==1:return httpx.Response(503)
        assert json.loads(request.content)['query']=='company'
        return httpx.Response(200,json={'results':[{'url':'https://example.com','title':'company'}]})
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as c:
        web=WebAdapter(HTTPWebSearchProvider('https://search.test',client=c),Store(tmp_path/'db'))
        assert await web.search('company')==await web.search('company')
        assert calls==2
    with pytest.raises(DatasourceError):await HTTPWebSearchProvider('').search('c')
    with pytest.raises(DatasourceError):await validate_public_url('file:///etc/passwd')
    with pytest.raises(DatasourceError):await validate_public_url('http://127.0.0.1')


def test_source_domain_config_blocks_previews_and_forums_and_ranks_official():
    from company_research.config import Settings
    from company_research.datasources import SourcePolicy
    policy=SourcePolicy.load(Settings(_env_file=None).source_domains)
    for url in ['https://www.renrendoc.com/paper/1.html','https://max.book118.com/html/2024/1.shtm',
                'https://www.doc88.com/p-1.html','https://www.docin.com/p-1.html',
                'https://guba.eastmoney.com/news,600519,1.html','https://xueqiu.com/1234567/890123']:
        assert policy.blocked(url),url
    for url in ['https://xueqiu.com/S/SH600519','https://www.stats.gov.cn/sj/','https://example.com/a']:
        assert not policy.blocked(url),url
    assert policy.rank('https://www.stats.gov.cn/x')<policy.rank('https://www.gov.cn/x')<policy.rank('https://example.com')
    assert policy.policy_site=='gov.cn'
    assert not SourcePolicy.load(Path('missing.yaml')).blocked('https://www.renrendoc.com/a')


async def test_filtered_results_do_not_use_the_read_budget():
    from company_research.datasources import SourcePolicy,UnifiedDatasource
    from company_research.schemas import SourceRecord
    urls=['https://www.renrendoc.com/1','https://max.book118.com/2','https://guba.eastmoney.com/3',
          'https://a.example.com/4','https://www.stats.gov.cn/5','https://b.example.com/6']
    class Web:
        def __init__(self):self.read_urls=[]
        async def search(self,query):return [{'url':u,'title':u} for u in urls]
        async def read(self,result):
            self.read_urls.append(result['url'])
            return [SourceRecord(source_id=result['url'],kind='web',title=result['url'],content=result['url'])]
    local=type('Local',(),{'search':lambda self,q:[],'errors':[]})()
    policy=SourcePolicy(['renrendoc.com','book118.com','guba.eastmoney.com'],[],['stats.gov.cn'])
    web=Web();found,_=await UnifiedDatasource(local,web,policy).retrieve('q',web_limit=2,filter_sources=True)
    assert web.read_urls==['https://www.stats.gov.cn/5','https://a.example.com/4']
    assert {s.source_id for s in found}==set(web.read_urls)
    # Without the flag (company/business searches) results are read in engine order.
    web=Web();await UnifiedDatasource(local,web,policy).retrieve('q',web_limit=2)
    assert web.read_urls==urls[:2]
