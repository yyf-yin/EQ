"""v1.1 acceptance: contracts, hard gates, bounds, restatement and local reuse."""
import asyncio
import json
from decimal import Decimal as D
from zipfile import ZipFile
import httpx
import pytest
from pypdf import PdfReader
from company_research.schemas import FinancialFact,ResearchBrief,SourceRecord,ResearchRun,ResearchTask
from company_research.framework import skeleton,validate_plan,descendants,CORE,coverage_matrix
from company_research.comparisons import compare,build_comparisons,core_gate,narrative
from company_research.contexts import extraction_batches,block_context,estimate_tokens
from company_research.llm import Gateway,ModelConfig,TruncatedOutput,IncompleteOutput,BudgetError,RoleLimits
from company_research.config import Settings
from company_research.storage import Store
from company_research.runner import Runner
from company_research.graphs import ResearchWorkflow
from company_research.researcher import researcher_graph
from company_research.financials.formulas import FinancialError,FinancialDataProcessor
from company_research.writers import report_snapshot


def brief(segments=None):
    return ResearchBrief(company_id='D',scope='fixture',historical_periods=['2023A','2024A','2025A'],
        forecast_periods=['2026E','2027E','2028E'],segments=segments or ['甲','乙'],interim_period='2026H1')


def fact(value,period,metric='parent_net_income',unit='CNY_million',**kwargs):
    return FinancialFact(company_id='D',metric_code=metric,period=period,unit=unit,value=D(str(value)),actual_or_forecast='actual',
        source_id='src-'+period,evidence_id='ev-'+period,locator='table row',validation_status='verified',**kwargs)


@pytest.mark.parametrize('segments,count',[(['甲','乙'],13),(['甲','乙','丙'],15)])
def test_framework_counts(segments,count):
    b=brief(segments);p=skeleton(b)
    assert len(p.blocks)==len(p.tasks)==count
    assert not validate_plan(p,b)
    assert set(CORE)=={'revenue','cost','parent_net_income','adjusted_parent_net_income','operating_cashflow'}


@pytest.mark.parametrize('fault',['missing','semantic','business','cycle','task','budget'])
def test_invalid_plan_detected(fault):
    b=brief();p=skeleton(b)
    if fault=='missing':p.blocks.pop()
    elif fault=='semantic':p.blocks[1].questions.pop(next(iter(p.blocks[1].questions)))
    elif fault=='business':p.blocks[6].segment_id='错误业务'
    elif fault=='cycle':p.blocks[0].dependencies=['X01']
    elif fault=='task':p.tasks.pop()
    elif fault=='budget':p.blocks[0].call_budget=9999
    assert validate_plan(p,b)


@pytest.mark.parametrize('previous,current,delta,rate,status',[
 (500,550,50,'.1','同比增加'),(100,80,-20,'-.2','同比减少'),(40,-10,-50,'-1.25','转亏'),
 (-100,-60,40,None,'亏损收窄'),(-100,-140,-40,None,'亏损扩大'),(-100,10,110,None,'扭亏'),
 (0,10,10,None,'不适用（同期基数为零）')])
def test_comparison_math(previous,current,delta,rate,status):
    c,r=compare(fact(current,'2025A'),fact(previous,'2024A'))
    assert c.delta==D(delta) and c.rate==(D(rate) if rate else None) and c.status==status
    assert all(x.input_refs for x in r)


def test_percentage_points_cash_and_periods():
    c,_=compare(fact('.32','2025A','gross_margin','ratio'),fact('.30','2024A','gross_margin','ratio'))
    assert c.delta==D(2) and c.rate is None
    c,_=compare(fact(-60,'2026H1','operating_cashflow'),fact(-100,'2025H1','operating_cashflow'))
    assert c.status=='净流出收窄' and c.delta==40
    for p in ['2025A','2025Q2','2025-12-31']:
        with pytest.raises(FinancialError):compare(fact(10,'2026H1'),fact(5,p))
    with pytest.raises(FinancialError):compare(fact(10,'2026Q2'),fact(5,'2025H1'))
    with pytest.raises(FinancialError):compare(fact(10,'2026H1'),fact(5,'2025H1',scope='parent'))


def test_long_material_windows_keep_comparison_headers():
    text='无关的文档开头\n'*3000+'单位：人民币百万元 合并\n| 指标 | 2025A | 2024A |\n|---|---|---|\n'+''.join(f'| 营业收入分类{i} | 550 | 500 |\n' for i in range(80))+'注：同期数据已经重述。\n'
    source=SourceRecord(kind='local',title='长年报',content=text,locator='p20 table3')
    task=ResearchTask(company_id='D',block_id='F01',topic='收入同比',role='financial',queries=['营业收入'])
    batches=extraction_batches([source],task.model_dump(mode='json'),['2025A','2024A'],1500)
    assert len(batches)>1
    for b in batches:
        assert any('2024A' in w['content'] and '百万元' in w['content'] and '重述' in w['content'] for w in b.windows)
        assert all(w['start']>len('无关的文档开头\n')*2900 for w in b.windows)


@pytest.fixture(scope='module')
def completed_v11(tmp_path_factory):
    path=tmp_path_factory.mktemp('v11-run')
    settings=Settings(database=path/'db',artifacts=path/'out',mock_delay=0)
    store=Store(settings.database);run=ResearchRun(company_input='v1.1 fixture',mode='mock');store.save_run(run)
    asyncio.run(Runner(settings,store).run(run.run_id))
    assert store.run(run.run_id).status=='limited',store.run(run.run_id).error
    return settings,store,run.run_id,store.report(run.run_id)


@pytest.mark.parametrize('fault',['baseline','parent_scope','prose','table','rate','interim','missing_comparison'])
def test_core_gate_faults(completed_v11,fault):
    report=completed_v11[3].model_copy(deep=True)
    if fault=='baseline':report.financial_facts=[f for f in report.financial_facts if f.period!='2022A']
    elif fault=='parent_scope':
        for f in report.financial_facts:
            if f.metric_code=='parent_net_income':f.scope='parent'
    elif fault=='prose':report.sections[3].children[0].paragraphs=[]
    elif fault=='table':report.sections[3].children[0].tables[0].rows[0][2]='999'
    elif fault=='rate':report.comparisons[0].rate=D(999)
    elif fault=='interim':report.comparisons[-1].previous_period='2025A'
    elif fault=='missing_comparison':report.comparisons.pop()
    issues=core_gate(report)
    assert issues and all(i.severity=='critical' and i.block_id and i.rule_id for i in issues)


def test_consistent_snapshot_and_exports(completed_v11):
    settings,store,id,report=completed_v11
    assert len(report.blocks)==13 and len(report.comparisons)==20 and not core_gate(report)
    assert report_snapshot(report)==report.approved_snapshot_id
    assert not any(e.status=='ready' for e in report.coverage)
    docx=settings.artifacts/id/'report.docx';pdf=settings.artifacts/id/'report.pdf'
    from docx import Document
    doc=Document(docx);doc_text=' '.join(p.text for p in doc.paragraphs)
    pdf_text=''.join(p.extract_text() for p in PdfReader(pdf).pages)
    for c in report.comparisons:
        assert narrative(c) in doc_text
        assert c.comparison_id in pdf_text
    changed=report.model_copy(deep=True);changed.sections[3].children[0].paragraphs.append('变更')
    assert report_snapshot(changed)!=changed.approved_snapshot_id


async def test_local_prose_repair_and_stale_approval(completed_v11):
    settings,store,id,report=completed_v11
    c=store.load_checkpoint(id)[1];c['report']=report.model_dump(mode='json')
    c['report']['sections'][3]['children'][0]['paragraphs']=[]
    workflow=ResearchWorkflow(settings,store,id)
    c.update(await workflow.final_review(c));assert c['issues']
    previous=c['report']['sections'][4]
    c.update(await workflow.revise(c));c.update(await workflow.final_review(c))
    assert not c['issues'] and c['report']['sections'][4]==previous
    c['report']['sections'][4]['paragraphs'].append('修改后未审核')
    with pytest.raises(FinancialError):await workflow.render(c)


def test_restatement_and_dag_reuse(completed_v11):
    _,store,id,report=completed_v11
    original=next(f for f in report.financial_facts if f.metric_code=='parent_net_income' and f.period=='2024A')
    updated=original.model_copy(update={'fact_id':'restated','value':original.value+10,'version':2,'supersedes':original.fact_id})
    found=FinancialDataProcessor().get(report.financial_facts+[updated],'DEMO','parent_net_income','2024A')
    assert found.fact_id=='restated'
    comparisons,_,_=build_comparisons(report.brief,report.financial_facts+[updated])
    assert [c for c in comparisons if c.current_period=='2025A' and c.metric_code=='parent_net_income'][0].previous_fact_id=='restated'
    affected=descendants([b.model_dump(mode='json') for b in report.blocks],{'B01-0'})
    assert {'B01-0','F04','V01','R01','X01'}.issubset(affected)
    assert not {'B01-1','I01-1','F01','F02'} & affected
    c=store.load_checkpoint(id)[1]
    before=block_context(c,'F01','analysis').snapshot_id
    updated_c=c|{'facts':[f|{'value':'999'} if f['segment_id']=='设备' else f for f in c['facts']]}
    assert block_context(updated_c,'F01','analysis').snapshot_id==before
    assert block_context(updated_c,'F04','analysis').snapshot_id!=block_context(c,'F04','analysis').snapshot_id


@pytest.mark.parametrize('protocol',['openai','anthropic'])
async def test_truncation_incomplete_and_budget(protocol,tmp_path,monkeypatch):
    monkeypatch.setenv('V11_TEST_KEY','test-only')
    calls=[]
    def handler(request):
        calls.append(request)
        if protocol=='openai':return httpx.Response(200,json={'choices':[{'message':{'tool_calls':[{'id':'o','function':{'name':'structured_output','arguments':'{'}}]},'finish_reason':'length'}],'usage':{'prompt_tokens':3,'completion_tokens':4}})
        return httpx.Response(200,json={'content':[],'stop_reason':'max_tokens','usage':{'input_tokens':3,'output_tokens':4}})
    config=ModelConfig(protocol=protocol,base_url='https://fixture.test/v1',model='fixture',api_key_env='V11_TEST_KEY')
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        gateway=Gateway({'test':config},{'default':'test'},client=client,global_call_budget=1,enforce_call_budgets=True)
        store=Store(tmp_path/'db');gateway.attach_store(store,'r')
        with pytest.raises(TruncatedOutput):await gateway.complete('researcher',[],schema=ResearchBrief)
        assert len(calls)==1 and gateway.call_records[0]['finish_reason'] in {'length','max_tokens'}
        with pytest.raises(BudgetError):await gateway.complete('researcher',[])
        assert len(calls)==1
        resumed=Gateway({'test':config},{'default':'test'},client=client,global_call_budget=1,enforce_call_budgets=True);resumed.attach_store(store,'r')
        with pytest.raises(BudgetError):await resumed.complete('researcher',[])


async def test_split_recovery_and_validated_batch_reuse(tmp_path):
    from company_research.agents import ResearchExtraction,CompressedResearch
    class Stub:
        def __init__(self):self.calls=[]
        def limits(self,role):return RoleLimits()
        async def structured(self,role,messages,schema):
            data=json.loads(messages[-1]['content']);self.calls.append((role,data))
            if role=='researcher' and '-' not in data['batch_id'][30:]:raise TruncatedOutput('injected')
            if role=='researcher':return ResearchExtraction(evidence=[],facts=[],followup_queries=[],gaps=[])
            return CompressedResearch(summary='test',evidence_ids=[],unresolved=[])
    source=SourceRecord(kind='local',title='test',content='单位：百万元\n营业收入 550 同期500\n'*60)
    class Datasource:
        async def retrieve(self,q):return [source],[]
    gateway=Stub();store=Store(tmp_path/'db');settings=Settings(extraction_input_tokens=8000,max_batch_splits=2)
    task=ResearchTask(company_id='D',block_id='F01',topic='收入',role='financial',queries=['营业收入']).model_dump(mode='json')
    state={'task':task,'brief':brief().model_dump(mode='json'),'round':0,'sources':[],'evidence':[],'facts':[],'queries':task['queries'],'gaps':[],'warnings':[],'summary':{}}
    graph=researcher_graph(gateway,Datasource(),store,'r',1,settings=settings)
    await graph.ainvoke(state)
    batches=store.records('r','batch')
    assert any(b['status']=='split' for b in batches) and any(b['parent_id'] and b['status']=='validated' for b in batches)
    assert all(b['company_id']=='D' and {'2022A','2025H1','2026H1'}.issubset(b['periods']) for b in batches)
    before=sum(role=='researcher' for role,_ in gateway.calls)
    await graph.ainvoke(state)
    after=sum(role=='researcher' for role,_ in gateway.calls)
    # Known split parents and validated child outputs are both reused.
    assert after==before


async def test_unrepairable_core_blocks_delivery(tmp_path,monkeypatch):
    original=ResearchWorkflow.resolve_company
    async def missing_baseline(self,c):
        result=await original(self,c)
        result['facts']=[f for f in result['facts'] if f['period']!='2022A']
        return result
    monkeypatch.setattr(ResearchWorkflow,'resolve_company',missing_baseline)
    settings=Settings(database=tmp_path/'db',artifacts=tmp_path/'out',mock_delay=0,max_replans=1,max_revisions=1)
    store=Store(settings.database);run=ResearchRun(company_input='core gap',mode='mock');store.save_run(run)
    await Runner(settings,store).run(run.run_id)
    assert store.run(run.run_id).status=='waiting_human'
    assert not (settings.artifacts/run.run_id/'report.docx').exists()
    assert store.records(run.run_id,'human_issue')
    assert not store.records(run.run_id,'calculation')


def test_legacy_reports_and_checkpoint_are_preserved(tmp_path):
    from company_research.reports import mock_report
    store=Store(tmp_path/'db');run=ResearchRun(company_input='old',mode='mock');store.save_run(run)
    report=mock_report('legacy');store.save_report(run.run_id,report)
    assert store.report(run.run_id).rules_version=='1.0'
    store.checkpoint(run.run_id,'write_sections',{'report':report.model_dump(mode='json')})
    with pytest.raises(FinancialError,match='旧检查点'):ResearchWorkflow(Settings(database=tmp_path/'db'),store,run.run_id)
    assert store.report(run.run_id).title==report.title


async def test_updated_fact_rejects_old_approved_snapshot(completed_v11):
    settings,store,id,report=completed_v11
    c=store.load_checkpoint(id)[1];c['report']=report.model_dump(mode='json')
    c['facts'][0]['value']='9999'
    with pytest.raises(FinancialError,match='事实已更新'):await ResearchWorkflow(settings,store,id).render(c)


async def test_business_restatement_reanalyzes_only_dependents(completed_v11):
    settings,store,id,_=completed_v11
    c=store.load_checkpoint(id)[1]
    workflow=ResearchWorkflow(settings,store,id)
    before=len(store.events(id))
    c.update(await workflow.analyze_blocks(c))
    assert not [e for e in store.events(id)[before:] if e.stage=='分块分析']
    original=next(f for f in c['facts'] if f['segment_id']=='设备' and f['metric_code']=='revenue' and f['period']=='2025A')
    # A restatement must have located directional evidence, not just supersedes.
    value=str(D(original['value'])+10)
    quote=f"{c['company']['name']} 2025年度 设备 营业收入由{original['value']}百万元更正为{value}百万元。"
    source=SourceRecord(source_id='restatement-source',kind='web',title=c['company']['name'],source_url='https://cninfo.com.cn/correction',content=quote).model_dump(mode='json')
    evidence={'evidence_id':'restatement-evidence','source_id':source['source_id'],'claim':quote,'quote':quote,'locator':'p1','actual_or_forecast':'actual','validation_status':'verified'}
    c['sources'].append(source);c['evidence'].append(evidence)
    c['facts'].append(original|{'fact_id':'restated-equipment-revenue','value':value,'version':2,'supersedes':original['fact_id'],'source_id':source['source_id'],'evidence_id':evidence['evidence_id']})
    workflow.bind(c)
    c.update(await workflow.financial_analysis(c));c.update(await workflow.assumptions(c));c.update(await workflow.calculate(c))
    before=len(store.events(id));c.update(await workflow.analyze_blocks(c))
    rerun={e.task_name for e in store.events(id)[before:] if e.stage=='分块分析' and e.status=='running'}
    assert rerun=={'B01-0','F04','V01','R01'}
    assert c['block_analyses']['F04']['version']!=completed_v11[3].blocks[4].version


def test_scope_and_assumption_versions_invalidate_context(completed_v11):
    c=completed_v11[1].load_checkpoint(completed_v11[2])[1]
    before=block_context(c,'F01','analysis').snapshot_id
    modified=c|{'brief':c['brief']|{'interim_period':'2026Q3'}}
    assert block_context(modified,'F01','analysis').snapshot_id!=before
    before=block_context(c,'V01','analysis').snapshot_id
    changed=[a|{'value':'0.15'} if i==0 else a for i,a in enumerate(c['assumptions'])]
    assert block_context(c|{'assumptions':changed},'V01','analysis').snapshot_id!=before


def test_cross_page_header_and_full_source():
    from company_research.datasources import join_document_chunks
    pages=[SourceRecord(kind='local',title='report.pdf',local_path='report.pdf',locator='page:1',content='合并 单位：百万元\n| 指标 | 2025A | 2024A |\n'),
           SourceRecord(kind='local',title='report.pdf',local_path='report.pdf',locator='page:2',content='| 营业收入 | 550 | 500 |\n注：比较数据为重述值。')]
    joined=join_document_chunks(pages)
    assert len(joined)==1 and '[page:1]' in joined[0].content and '[page:2]' in joined[0].content
    t=ResearchTask(company_id='D',block_id='F01',topic='收入',role='financial',queries=['营业收入'])
    b=extraction_batches(joined,t.model_dump(mode='json'),['2025A'],5000)
    assert b and '2024A' in b[0].windows[0]['content'] and '重述' in b[0].windows[0]['content']


def test_parent_attribution_cannot_be_inferred():
    from company_research.agents import ResearchExtraction,FactProposal,bind_extraction
    source=SourceRecord(kind='local',title='wrong_scope',content='2025年合并净利润 100 百万元。')
    output=ResearchExtraction(evidence=[],facts=[FactProposal(source_id=source.source_id,quote=source.content,metric_code='parent_net_income',period='2025A',unit='CNY_million',value=100,actual_or_forecast='actual')],gaps=[],followup_queries=[])
    _,facts,warnings=bind_extraction(output,[source],'D')
    assert not facts and any('归母口径' in w for w in warnings)


async def test_small_role_capacity_splits_and_records_actual_content(tmp_path,monkeypatch):
    from company_research.agents import ask_bounded,FinancialAnalysis
    monkeypatch.setenv('V11_TEST_KEY','fixture-only');seen=[]
    def handler(request):
        data=json.loads(request.content);seen.append(data)
        value={'observations':['only selected facts'],'evidence_ids':[],'calculation_ids':[],'limitations':[],'question_answers':{'F01:q1':'selected data'}}
        return httpx.Response(200,json={'choices':[{'message':{'tool_calls':[{'id':'o','function':{'name':'structured_output','arguments':json.dumps(value)}}]},'finish_reason':'tool_calls'}],'usage':{'prompt_tokens':20,'completion_tokens':30}})
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        gateway=Gateway({'t':ModelConfig(protocol='openai',base_url='https://fixture.test/v1',model='fixture',api_key_env='V11_TEST_KEY',max_tokens=1000)},
            {'default':'t'},client=client,role_limits={'default':{'input_capacity':6000,'context_capacity':9000,'output_reserve':1000}},global_token_budget=100000)
        store=Store(tmp_path/'db');gateway.attach_store(store,'r')
        context={'block':{'block_id':'F01'},'facts':[{'fact_id':str(i),'value':str(i),'description':'x'*900} for i in range(20)],'evidence':[],'calculations':[]}
        result=await ask_bounded(gateway,'financial_analyzer',FinancialAnalysis,context)
        assert result.observations and len(seen)>1
        assert all(estimate_tokens(p)<=6000 for p in seen)
        assert any(c['content'].get('partial_context') for c in store.records('r','sent_context'))


async def test_role_concurrency_bounds_parallel_requests(monkeypatch):
    monkeypatch.setenv('V11_TEST_KEY','fixture-only');active=peak=0
    async def handler(request):
        nonlocal active,peak
        active+=1;peak=max(peak,active);await asyncio.sleep(.02);active-=1
        return httpx.Response(200,json={'choices':[{'message':{'content':'ok'},'finish_reason':'stop'}]})
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        gateway=Gateway({'t':ModelConfig(protocol='openai',base_url='https://fixture.test/v1',model='fixture',api_key_env='V11_TEST_KEY')},
            {'default':'t'},client=client,concurrency=3,role_limits={'researcher':{'concurrency':1}})
        await asyncio.gather(*(gateway.complete('researcher',[{'role':'user','content':'fixture'}]) for _ in range(3)))
        assert peak==1 and gateway.calls==3


def test_configurable_cache_freshness(tmp_path):
    store=Store(tmp_path/'db');run=ResearchRun(company_input='freshness',mode='mock');store.save_run(run)
    store.cache('search-disclosure',{'value':'old'});assert store.cached('search-disclosure')
    ResearchWorkflow(Settings(database=tmp_path/'db',cache_ttl_hours=0),store,run.run_id)
    assert store.cached('search-disclosure') is None


